import json
import os
import sys
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter


OUTPUT_PATH = sys.argv[1]
PAYLOAD = json.load(sys.stdin)
BR_TZ = ZoneInfo("America/Sao_Paulo")

NAVY = "0B1F3A"
BLUE = "1261A0"
TEAL = "009E8E"
LIGHT_BLUE = "DCEEFF"
WHITE = "FFFFFF"
GRAY = "F3F6FA"
GREEN = "DDF4E7"
YELLOW = "FFF3CD"


def to_brt(value):
    if not value:
        return ""
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone(BR_TZ).replace(tzinfo=None)
        return parsed
    except (TypeError, ValueError):
        return str(value)


def clean(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    return str(value).strip()


def add_table(ws, name, headers, rows, tab_color):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab_color
    ws.freeze_panes = "A2"
    ws.append(headers)

    header_fill = PatternFill("solid", fgColor=NAVY)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ws.append(row)

    ws.auto_filter.ref = ws.dimensions
    if rows:
        table = Table(displayName=name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm"

    for index, header in enumerate(headers, start=1):
        max_length = len(header)
        for row in ws.iter_rows(min_row=2, min_col=index, max_col=index):
            value = row[0].value
            if value is not None:
                max_length = max(max_length, len(str(value)))
        ws.column_dimensions[get_column_letter(index)].width = min(max(max_length + 2, 13), 38)

    for row_index in range(2, ws.max_row + 1):
        if row_index % 2 == 0:
            for cell in ws[row_index]:
                cell.fill = PatternFill("solid", fgColor=GRAY)
    ws.row_dimensions[1].height = 32


def set_column_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


workbook = Workbook()
summary_sheet = workbook.active
summary_sheet.title = "Resumo"
summary_sheet.sheet_view.showGridLines = False
summary_sheet.sheet_properties.tabColor = TEAL

summary_sheet.merge_cells("A1:D1")
summary_sheet["A1"] = "Cadastro completo de atletas — /argentina"
summary_sheet["A1"].font = Font(color=WHITE, bold=True, size=16)
summary_sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
summary_sheet["A1"].alignment = Alignment(horizontal="center")
summary_sheet.row_dimensions[1].height = 30

summary_rows = [
    ("Gerado em", to_brt(PAYLOAD.get("generated_at"))),
    ("Página incluída", "/argentina"),
    ("Cadastros diretos", PAYLOAD["summary"]["total_registrations"]),
    ("Comprovantes recebidos", PAYLOAD["summary"]["with_receipt"]),
    ("Pendentes de pagamento", PAYLOAD["summary"]["pending_payment"]),
    ("Pagamentos validados", PAYLOAD["summary"]["validated_payment"]),
    ("Possíveis duplicidades para revisar", PAYLOAD["summary"]["possible_duplicates"]),
]

for label, value in summary_rows:
    summary_sheet.append([label, value])

for row in summary_sheet.iter_rows(min_row=2, max_col=2):
    row[0].font = Font(bold=True, color=NAVY)
    row[0].fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    row[1].alignment = Alignment(wrap_text=True)
    if isinstance(row[1].value, datetime):
        row[1].number_format = "dd/mm/yyyy hh:mm"

summary_sheet["A11"] = "Observação"
summary_sheet["A11"].font = Font(bold=True, color=NAVY)
summary_sheet["B11"] = "A aba principal mantém todos os registros sem excluir linhas. O indicador de duplicidade considera e-mail e telefone iguais e deve ser revisado antes de qualquer remoção."
summary_sheet["B11"].alignment = Alignment(wrap_text=True, vertical="top")
summary_sheet["B11"].fill = PatternFill("solid", fgColor=YELLOW)
summary_sheet.row_dimensions[11].height = 48
summary_sheet["A13"] = "Privacidade"
summary_sheet["A13"].font = Font(bold=True, color=NAVY)
summary_sheet["B13"] = "Esta planilha contém dados pessoais de atletas e responsáveis. Compartilhe somente com a equipe autorizada."
summary_sheet["B13"].alignment = Alignment(wrap_text=True, vertical="top")
summary_sheet["B13"].fill = PatternFill("solid", fgColor=GREEN)
summary_sheet.row_dimensions[13].height = 36
set_column_widths(summary_sheet, {1: 34, 2: 76, 3: 3, 4: 3})

registration_headers = [
    "ID do cadastro",
    "Data do cadastro",
    "Nome do atleta",
    "Idade",
    "Posição",
    "Telefone do atleta",
    "E-mail",
    "Responsável",
    "Telefone do responsável",
    "Consentimento",
    "Consentimento em",
    "Status do pagamento",
    "Valor",
    "Moeda",
    "Método",
    "Alias",
    "CVU",
    "Titular",
    "Comprovante recebido",
    "Status do comprovante",
    "Referência do comprovante",
    "Enviado em",
    "Campanha",
    "Datas do evento",
    "UTM origem",
    "UTM mídia",
    "UTM campanha",
    "UTM anúncio",
    "UTM termo",
    "URL de origem",
    "URL de destino",
    "Referência",
    "ID do evento Meta",
    "Possível duplicidade",
    "Atualizado em",
]

registration_rows = []
for item in PAYLOAD["registrations"]:
    registration_rows.append([
        item["registration_id"],
        to_brt(item["registered_at"]),
        item["athlete_name"],
        item["athlete_age"],
        item["position"],
        item["phone"],
        item["email"],
        item["guardian_name"],
        item["guardian_phone"],
        item["consent"],
        to_brt(item["consent_at"]),
        item["payment_status"],
        item["payment_amount"],
        item["payment_currency"],
        item["payment_method"],
        item["payment_alias"],
        item["payment_cvu"],
        item["payment_holder"],
        item["receipt_received"],
        item["receipt_status"],
        item["receipt_reference"],
        to_brt(item["receipt_uploaded_at"]),
        item["campaign"],
        item["event_dates"],
        item["utm_source"],
        item["utm_medium"],
        item["utm_campaign"],
        item["utm_content"],
        item["utm_term"],
        item["source_url"],
        item["landing_url"],
        item["referrer"],
        item["meta_event_id"],
        item["possible_duplicate"],
        to_brt(item["updated_at"]),
    ])

registrations_sheet = workbook.create_sheet("Cadastros completos")
add_table(registrations_sheet, "CadastrosArgentina", registration_headers, registration_rows, BLUE)
set_column_widths(registrations_sheet, {
    1: 38, 2: 18, 3: 30, 4: 10, 5: 20, 6: 20, 7: 34, 8: 30, 9: 22,
    10: 14, 11: 18, 12: 25, 13: 16, 14: 12, 15: 18, 16: 26, 17: 28, 18: 30,
    19: 21, 20: 25, 21: 25, 22: 18, 23: 34, 24: 20, 25: 18, 26: 18, 27: 32,
    28: 36, 29: 24, 30: 34, 31: 48, 32: 42, 33: 40, 34: 24, 35: 18,
})

receipt_headers = [
    "ID do comprovante",
    "ID do cadastro vinculado",
    "Enviado em",
    "Atleta",
    "Telefone",
    "E-mail",
    "Status do pagamento",
    "Status do comprovante",
    "Referência",
    "Arquivo",
    "Tipo",
    "Tamanho",
    "Valor",
]

receipt_rows = [
    [
        item["receipt_id"],
        item["linked_registration_id"],
        to_brt(item["uploaded_at"]),
        item["athlete_name"],
        item["phone"],
        item["email"],
        item["payment_status"],
        item["receipt_status"],
        item["receipt_reference"],
        item["file_name"],
        item["file_type"],
        item["file_size"],
        item["payment_amount"],
    ]
    for item in PAYLOAD["receipts"]
]

receipts_sheet = workbook.create_sheet("Comprovantes")
add_table(receipts_sheet, "ComprovantesArgentina", receipt_headers, receipt_rows, TEAL)
set_column_widths(receipts_sheet, {1: 38, 2: 38, 3: 18, 4: 30, 5: 20, 6: 34, 7: 25, 8: 25, 9: 26, 10: 38, 11: 24, 12: 16, 13: 16})

workbook.properties.title = "Cadastro completo de atletas — Argentina"
workbook.properties.subject = "Cadastros diretos da página /argentina"
workbook.properties.creator = "EC10 Talentos"
os.makedirs(os.path.dirname(OUTPUT_PATH), exist_ok=True)
workbook.save(OUTPUT_PATH)
